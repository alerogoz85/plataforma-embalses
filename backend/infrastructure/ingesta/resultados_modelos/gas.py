"""Lee las salidas del modelo de gas y arma un ResultadoModelo.

Archivos (bajo la carpeta del proyecto, en `Resultados/`):
* pronostico_v2/Pronostico_Gas_v2.xlsx (script 4 v2): pronostico coherente por campo y nacional, con backtest y referencia de Arps
* pronostico_dca/Pronostico_DCA_Declaraciones.xlsx (script 8): curvas de declinacion de Arps por cohorte de declaracion (opcional)
* backtesting/*.csv y comparacion/resumen_comparacion_por_anio.csv: error de las declaraciones frente al gas observado (opcionales)

Hace falta al menos el libro v2 o el de curvas de declinacion. No se recalcula nada: se muestran las tablas y series tal como las dejo el
modelo. El libro v2 trae 16 series (el total nacional y 15 campos) x 3 variables; por eso el resultado declara dos DIMENSIONES de
filtro (serie y variable) y un grafico por combinacion."""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook

from domain.entities.modelo_analitico import (
    DimensionFiltro,
    GraficoResultado,
    Indicador,
    LineaGrafico,
    PuntoSerie,
    ResultadoModelo,
    TablaResultado,
)
from infrastructure.ingesta.resultados_modelos.comun import (
    TOPE_FILAS_TABLA,
    SalidaModeloInvalidaError,
    celda,
    mes_iso,
    momento_de,
    numero,
)

RUTA_V2 = Path("Resultados") / "pronostico_v2" / "Pronostico_Gas_v2.xlsx"
RUTA_DCA = Path("Resultados") / "pronostico_dca" / "Pronostico_DCA_Declaraciones.xlsx"
CARPETA_BACKTEST = Path("Resultados") / "backtesting"
RUTA_COMPARACION = Path("Resultados") / "comparacion" / "resumen_comparacion_por_anio.csv"

NACIONAL = "NACIONAL"
VARIABLES = ("Fiscalizado", "Comercializado", "Entregado a gasoducto")
MESES_HISTORIA_GRAFICO = 36  # el modelo ajusta Arps sobre los ultimos 36 meses: el grafico muestra la misma ventana
ETQ_MME = "Oferta gas declarado (MME)"
ETQ_ANH = "Oferta gas comercializado (ANH)"
CURVAS = ("Proyeccion Exponencial", "Proyeccion Hiperbolica", "Proyeccion Armonica")
HOJAS_V2 = {"Info", "Resumen", "Pronostico", "Proporciones", "Backtest", "Arps hiperbolica", "Declarado nacional", "Historico", "Arps mensual"}
HOJAS_DCA = {"Parametros", "Proyeccion mensual", "Resumen", "Metricas vs MME", "Metricas vs ANH real"}


# ----------------------------------------------------------------------------------------------- hojas
def _hoja(libro, nombre: str, archivo: str, columnas: set[str] | None = None) -> list[dict[str, Any]]:
    """Filas de una hoja como diccionarios por encabezado. Ignora filas vacias y las notas de pie (una sola celda con texto)."""
    if nombre not in libro.sheetnames:
        raise SalidaModeloInvalidaError(f"{archivo}: falta la hoja '{nombre}'")
    filas = list(libro[nombre].iter_rows(values_only=True))
    if not filas:
        raise SalidaModeloInvalidaError(f"{archivo}: la hoja '{nombre}' esta vacia")
    encabezado = [str(c).strip() if c is not None else "" for c in filas[0]]
    if columnas and not columnas <= set(encabezado):
        raise SalidaModeloInvalidaError(f"{archivo}: la hoja '{nombre}' no tiene las columnas {sorted(columnas - set(encabezado))}")
    salida = []
    for fila in filas[1:]:
        if sum(1 for c in fila if c is not None) <= 1:  # fila vacia o nota al pie
            continue
        salida.append({h: v for h, v in zip(encabezado, fila) if h})
    return salida


def _pares_info(libro) -> list[tuple[str, str]]:
    """Pares (etiqueta, texto) de la hoja Info; las lineas de continuacion (sin etiqueta) se unen a la anterior."""
    pares: list[tuple[str, str]] = []
    for fila in libro["Info"].iter_rows(values_only=True):
        a = str(fila[0]).strip() if fila[0] is not None else ""
        b = str(fila[1]).strip() if len(fila) > 1 and fila[1] is not None else ""
        if a and b:
            pares.append((a, b))
        elif b and pares:
            pares[-1] = (pares[-1][0], f"{pares[-1][1]} {b}")
    return pares


def _tabla(id_: str, titulo: str, filas: list[dict[str, Any]], nota: Optional[str] = None, filtros: tuple[tuple[str, str], ...] = ()) -> TablaResultado:
    columnas = tuple(filas[0].keys()) if filas else ("Sin filas",)
    recortada = len(filas) > TOPE_FILAS_TABLA
    cuerpo = tuple(tuple(celda(f.get(c)) for c in columnas) for f in filas[:TOPE_FILAS_TABLA]) if filas else ()
    if recortada:
        nota = (nota + " " if nota else "") + f"Se muestran las primeras {TOPE_FILAS_TABLA} de {len(filas)} filas."
    return TablaResultado(id_, titulo, columnas, cuerpo, nota, tuple((d, c) for d, c in filtros if c in columnas))


def _tabla_csv(id_: str, titulo: str, ruta: Path, nota: Optional[str] = None, renombrar: Optional[dict[str, str]] = None) -> TablaResultado:
    with ruta.open(encoding="utf-8-sig", newline="") as f:
        filas = list(csv.DictReader(f))
    if not filas:
        raise SalidaModeloInvalidaError(f"{ruta.name} no tiene filas")
    renombrar = renombrar or {}
    filas = [{renombrar.get(k or "", k or "Resumen"): _a_numero(v) for k, v in fila.items()} for fila in filas]
    return _tabla(id_, titulo, filas, nota)


def _a_numero(valor: Any) -> Any:
    try:
        return float(valor)
    except (TypeError, ValueError):
        return valor


# ----------------------------------------------------------------------------------------------- entrada
def leer_gas(carpeta: Path) -> ResultadoModelo:
    carpeta = Path(carpeta)
    v2, dca = carpeta / RUTA_V2, carpeta / RUTA_DCA
    if not v2.is_file() and not dca.is_file():
        raise SalidaModeloInvalidaError(f"no hay ningun libro de resultados: se buscaron {v2} y {dca}")

    dimensiones: list[DimensionFiltro] = []
    graficos: list[GraficoResultado] = []
    tablas: list[TablaResultado] = []
    indicadores: list[Indicador] = []
    advertencias: list[str] = []
    configuracion: list[tuple[str, str]] = []
    momentos, fuentes = [], []

    if v2.is_file():
        _leer_v2(v2, dimensiones, graficos, tablas, indicadores, advertencias, configuracion)
        momentos.append(momento_de(v2))
        fuentes.append("pronóstico coherente por campo (v2)")
    if dca.is_file():
        _leer_dca(dca, graficos, tablas, indicadores, advertencias, configuracion)
        momentos.append(momento_de(dca))
        fuentes.append("curvas de declinación de las declaraciones (Arps)")
    _leer_backtesting_declaraciones(carpeta, tablas, indicadores, fuentes, momentos)

    return ResultadoModelo(
        modelo_id="gas", generado_en=max(momentos), origen="Modelo de gas fiscalizado y declarado · " + " + ".join(fuentes),
        configuracion=tuple(configuracion), indicadores=tuple(indicadores), tablas=tuple(tablas), advertencias=tuple(dict.fromkeys(advertencias)),
        dimensiones=tuple(dimensiones), graficos=tuple(graficos),
    )


# ----------------------------------------------------------------------------------------------- libro v2
def _leer_v2(ruta: Path, dimensiones, graficos, tablas, indicadores, advertencias, configuracion) -> None:
    archivo = ruta.name
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        faltan = HOJAS_V2 - set(libro.sheetnames)
        if faltan:
            raise SalidaModeloInvalidaError(f"{archivo}: faltan hojas {sorted(faltan)}")
        resumen = _hoja(libro, "Resumen", archivo, {"Serie", "Variable"})
        pronostico = _hoja(libro, "Pronostico", archivo, {"Serie", "Variable", "Fecha", "Pronostico", "IC80 inferior", "IC80 superior"})
        historico = _hoja(libro, "Historico", archivo, {"Serie", "Variable", "Fecha", "Valor (MPCD)"})
        arps = _hoja(libro, "Arps mensual", archivo, {"Serie", "Variable", "Fecha", "Proyeccion Arps (MPCD)"})
        declarado = _hoja(libro, "Declarado nacional", archivo, {"Fecha", ETQ_MME})
        proporciones = _hoja(libro, "Proporciones", archivo)
        backtest = _hoja(libro, "Backtest", archivo)
        arps_param = _hoja(libro, "Arps hiperbolica", archivo)
        info = _pares_info(libro)
    finally:
        libro.close()
    if not pronostico:
        raise SalidaModeloInvalidaError(f"{archivo}: la hoja 'Pronostico' no tiene filas")

    series_ord = [NACIONAL] + sorted({f["Serie"] for f in pronostico} - {NACIONAL})
    variables = [v for v in VARIABLES if any(f["Variable"] == v for f in pronostico)]
    dimensiones.append(DimensionFiltro("serie", "Campo", tuple(series_ord), NACIONAL if NACIONAL in series_ord else series_ord[0]))
    dimensiones.append(DimensionFiltro("variable", "Variable", tuple(variables), variables[0]))

    def por_clave(filas, columna_valor):
        d: dict[tuple[str, str], list[tuple[str, float]]] = {}
        for f in filas:
            v = numero(f[columna_valor], archivo, columna_valor, opcional=True)
            if v is not None:
                d.setdefault((f["Serie"], f["Variable"]), []).append((mes_iso(str(f["Fecha"]), archivo), v))
        return d

    hist = por_clave(historico, "Valor (MPCD)")
    arps_ref = por_clave(arps, "Proyeccion Arps (MPCD)")
    pron: dict[tuple[str, str], list[PuntoSerie]] = {}
    for f in sorted(pronostico, key=lambda f: str(f["Fecha"])):
        y = numero(f["Pronostico"], archivo, "Pronostico")
        inf = numero(f["IC80 inferior"], archivo, "IC80 inferior", opcional=True)
        sup = numero(f["IC80 superior"], archivo, "IC80 superior", opcional=True)
        con_banda = inf is not None and sup is not None and inf <= y <= sup
        pron.setdefault((f["Serie"], f["Variable"]), []).append(
            PuntoSerie(mes_iso(str(f["Fecha"]), archivo), y, inf if con_banda else None, sup if con_banda else None)
        )
    mme = [(mes_iso(str(f["Fecha"]), archivo), numero(f[ETQ_MME], archivo, ETQ_MME)) for f in declarado if f.get(ETQ_MME) is not None]

    for serie in series_ord:
        for variable in variables:
            puntos_pron = tuple(pron.get((serie, variable), ()))
            if not puntos_pron:
                continue  # el modelo no pronostico esta combinacion (la advertencia de la tabla Resumen dice por que)
            lineas = []
            observado = sorted(hist.get((serie, variable), []))[-MESES_HISTORIA_GRAFICO:]
            if observado:
                lineas.append(LineaGrafico("Observado", tuple(PuntoSerie(x, y) for x, y in observado)))
            lineas.append(LineaGrafico("Pronóstico (banda: intervalo empírico 80 %)", puntos_pron))
            ref = [(x, y) for x, y in sorted(arps_ref.get((serie, variable), [])) if puntos_pron[0].x <= x <= puntos_pron[-1].x]
            if ref:
                lineas.append(LineaGrafico("Referencia: Arps hiperbólica", tuple(PuntoSerie(x, y) for x, y in ref), "punteada"))
            if serie == NACIONAL and variable == "Comercializado" and mme:
                desde = observado[0][0] if observado else puntos_pron[0].x
                visibles = [(x, y) for x, y in mme if desde <= x <= puntos_pron[-1].x]
                if visibles:
                    lineas.append(LineaGrafico("Oferta declarada (MME, declaración 2026)", tuple(PuntoSerie(x, y) for x, y in visibles), "punteada"))
            graficos.append(GraficoResultado(
                f"v2:{serie}:{variable}", f"{serie} · {variable}: observado y pronóstico", "MPCD", tuple(lineas),
                (("serie", serie), ("variable", variable)),
            ))

    # Tablas filtrables por campo y variable.
    filtros_sv = (("serie", "Serie"), ("variable", "Variable"))
    tablas.append(_tabla("v2_resumen", "Pronóstico a 12 meses y calidad del backtest", resumen,
                         "Variación % = promedio del pronóstico de 12 meses frente al de los últimos 12 meses reales. MASE < 1: mejor que el ingenuo a un paso.", filtros_sv))
    if backtest:
        tablas.append(_tabla("v2_backtest", "Backtest: método principal frente a pronosticar cada variable por separado", backtest,
                             "BIAS = pronóstico − real (positivo: sobreestima). «Principal = SI» es el método que usa el pronóstico.", filtros_sv))
    if proporciones:
        tablas.append(_tabla("v2_proporciones", "Proporciones que derivan las otras dos variables del fiscalizado", proporciones,
                             "p1 = comercializado / fiscalizado; p2 = entregado a gasoducto / comercializado (p2 ≤ 1).", (("serie", "Serie"),)))
    if arps_param:
        tablas.append(_tabla("v2_arps", "Curva de Arps hiperbólica (referencia de declinación)", arps_param,
                             "A nivel nacional los parámetros de Arps no tienen interpretación física.", filtros_sv))

    # Indicadores del total nacional.
    nac = {f["Variable"]: f for f in resumen if f["Serie"] == NACIONAL}
    base = nac.get("Fiscalizado")
    if base:
        prom = numero(base.get("Promedio pronostico 12m (MPCD)"), archivo, "Promedio pronostico 12m", opcional=True)
        var = numero(base.get("Variacion %"), archivo, "Variacion %", opcional=True)
        prom_hist = numero(base.get("Promedio ultimos 12m (MPCD)"), archivo, "Promedio ultimos 12m", opcional=True)
        if prom is not None and var is not None and prom_hist is not None:
            indicadores.append(Indicador(
                "Gas fiscalizado nacional, próximos 12 meses", f"{prom:,.0f} MPCD".replace(",", "."),
                f"{var:+.1f} % frente a los últimos 12 meses reales ({prom_hist:,.0f} MPCD)".replace(",", "."), None,
            ))
        mejora = numero(base.get("Mejora vs Naive %"), archivo, "Mejora vs Naive %", opcional=True)
        wape = numero(base.get("WAPE backtest %"), archivo, "WAPE backtest %", opcional=True)
        if mejora is not None and wape is not None:
            indicadores.append(Indicador("Mejora frente al pronóstico ingenuo (nacional)", f"{mejora:.0f} %", f"En el backtest · error relativo (WAPE) {wape:.1f} %", None))
    con_adv = [f for f in resumen if f.get("Advertencia")]
    if resumen:
        ingenuo = sum(1 for f in con_adv if "naive tuvo menor error" in str(f["Advertencia"]))
        indicadores.append(Indicador(
            "Series con advertencia del modelo", f"{len(con_adv)} de {len(resumen)}",
            f"En {ingenuo} el pronóstico ingenuo tuvo menor error en el backtest", "alerta" if con_adv else "ok",
        ))
    for etiqueta, texto in info:
        if etiqueta.startswith("Nota"):
            advertencias.append(texto)
        elif etiqueta in ("Insumo", "Ultimo mes con dato", "Horizonte", "Unidades", "Metodo", "Intervalo"):
            configuracion.append(({"Ultimo mes con dato": "Último mes con dato", "Metodo": "Método"}.get(etiqueta, etiqueta), texto))


# ----------------------------------------------------------------------------------------------- libro DCA
def _leer_dca(ruta: Path, graficos, tablas, indicadores, advertencias, configuracion) -> None:
    archivo = ruta.name
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        faltan = HOJAS_DCA - set(libro.sheetnames)
        if faltan:
            raise SalidaModeloInvalidaError(f"{archivo}: faltan hojas {sorted(faltan)}")
        resumen = _hoja(libro, "Resumen", archivo)
        mensual = _hoja(libro, "Proyeccion mensual", archivo)
        parametros = _hoja(libro, "Parametros", archivo)
        met_anh = _hoja(libro, "Metricas vs ANH real", archivo)
        met_mme = _hoja(libro, "Metricas vs MME", archivo)
        info = _pares_info(libro) if "Info" in libro.sheetnames else []
    finally:
        libro.close()
    if not mensual:
        raise SalidaModeloInvalidaError(f"{archivo}: la hoja 'Proyeccion mensual' no tiene filas")
    for columna in ("Cohorte", "Periodo", ETQ_ANH, ETQ_MME, *CURVAS):
        if columna not in mensual[0]:
            raise SalidaModeloInvalidaError(f"{archivo}: 'Proyeccion mensual' no tiene la columna '{columna}'")

    cohorte = max(int(numero(f["Cohorte"], archivo, "Cohorte")) for f in mensual)
    de_la_cohorte = sorted((f for f in mensual if int(numero(f["Cohorte"], archivo, "Cohorte")) == cohorte), key=lambda f: str(f["Periodo"]))
    lineas = []
    for columna, nombre, estilo in [(ETQ_ANH, "Gas comercializado observado (ANH)", "continua"), (ETQ_MME, "Oferta declarada (MME)", "continua"),
                                    ("Proyeccion Exponencial", "Curva exponencial", "punteada"), ("Proyeccion Hiperbolica", "Curva hiperbólica", "punteada"),
                                    ("Proyeccion Armonica", "Curva armónica", "punteada")]:
        puntos = tuple(
            PuntoSerie(mes_iso(str(f["Periodo"]), archivo), y)
            for f in de_la_cohorte
            if (y := numero(f[columna], archivo, columna, opcional=True)) is not None
        )
        if puntos:
            lineas.append(LineaGrafico(nombre, puntos, estilo))
    if lineas:
        graficos.append(GraficoResultado(f"dca:{cohorte}", f"Declaración {cohorte}: observado, declarado y curvas de declinación", "MPCD", tuple(lineas)))

    validables = [f for f in resumen if str(f.get("Validacion_ANH_real", "")).startswith(("DCA", "MME"))]
    dca_gana = sum(1 for f in validables if str(f["Validacion_ANH_real"]).startswith("DCA"))
    mapes = sorted(v for f in resumen if (v := numero(f.get("MAPE_%"), archivo, "MAPE_%", opcional=True)) is not None)
    indicadores.append(Indicador("Cohortes de declaración con curvas de declinación", str(len(resumen)), f"cohorte más reciente: {cohorte}", None))
    if validables:
        indicadores.append(Indicador("Cohortes donde la curva supera a la declaración", f"{dca_gana} de {len(validables)}",
                                     "frente al gas comercializado real fuera de la ventana de ajuste", None))
    if mapes:
        indicadores.append(Indicador("MAPE mediano del mejor modelo frente a la declaración", f"{mapes[len(mapes) // 2]:.1f} %", "por cohorte, sobre el segmento evaluado", None))

    tablas.append(_tabla("dca_resumen", "Curvas de declinación: resumen por cohorte de declaración", resumen))
    tablas.append(_tabla("dca_metricas_anh", "Error de cada curva y de la declaración frente al gas comercializado real", met_anh))
    tablas.append(_tabla("dca_metricas_mme", "Error de cada curva frente a la declaración", met_mme))
    tablas.append(_tabla("dca_parametros", "Parámetros ajustados de las curvas", parametros, "Di_mensual en 1/mes; en el agregado nacional b y Di no tienen interpretación física."))
    for etiqueta, texto in info:
        if etiqueta.startswith("Nota") and texto:
            advertencias.append(texto)


# ----------------------------------------------------------------------------------------------- backtesting de las declaraciones
def _leer_backtesting_declaraciones(carpeta: Path, tablas, indicadores, fuentes, momentos) -> None:
    base = carpeta / CARPETA_BACKTEST
    global_csv = base / "backtest_resumen_global.csv"
    por_anio = base / "backtest_por_anio_declaracion.csv"
    skill = base / "backtest_skill_vs_naive_horizonte.csv"
    comparacion = carpeta / RUTA_COMPARACION
    hubo = False
    if global_csv.is_file():
        with global_csv.open(encoding="utf-8-sig", newline="") as f:
            filas = list(csv.DictReader(f))
        if filas:
            g = filas[0]
            n = numero(g.get("n"), global_csv.name, "n", opcional=True)
            mape = numero(g.get("MAPE"), global_csv.name, "MAPE", opcional=True)
            sesgo = numero(g.get("sesgo_pct"), global_csv.name, "sesgo_pct", opcional=True)
            dentro = numero(g.get("pct_dentro_10"), global_csv.name, "pct_dentro_10", opcional=True)
            if mape is not None and n is not None:
                detalle = f"{n:.0f} comparaciones mensuales"
                if sesgo is not None:
                    detalle += f" · sesgo {sesgo:+.1f} %"
                if dentro is not None:
                    detalle += f" · {dentro:.0f} % dentro de ±10 %"
                indicadores.append(Indicador("Error medio de las declaraciones frente a lo comercializado (MAPE)", f"{mape:.1f} %", detalle, None))
            tablas.append(_tabla_csv("bt_global", "Backtesting de las declaraciones: resumen global", global_csv, renombrar={"": "Resumen"}))
            hubo = True
            momentos.append(momento_de(global_csv))
    if por_anio.is_file():
        tablas.append(_tabla_csv("bt_anio", "Backtesting de las declaraciones por año de declaración", por_anio,
                                 "sesgo = declaración menos lo comercializado; ratio = comercializado / declarado."))
        hubo = True
    if skill.is_file():
        tablas.append(_tabla_csv("bt_skill", "Declaración frente al pronóstico ingenuo, por horizonte (años)", skill,
                                 "skill < 0: la declaración tuvo mayor error que repetir el último valor observado."))
        hubo = True
    if comparacion.is_file():
        tablas.append(_tabla_csv("comparacion_anio", "Declarado frente a comercializado por año de declaración", comparacion))
        hubo = True
    if hubo:
        fuentes.append("backtesting de las declaraciones")
