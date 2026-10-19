"""Lectura y validacion de los archivos que produce el modelo de carbon termico.

El modelo vive en su propio proyecto (`Modelo de carbon termico`) y entrega:

* `resultados/proyeccion_carbon.xlsx`: proyeccion mensual, trimestres, inventario,
  comparacion con XM, validacion por episodios y supuestos.
* `resultados/evaluacion/evaluacion_modelo.xlsx`: veredicto de validez y error por horizonte.
* `datos/base/base_mensual_nacional.parquet`: historico mensual del consumo de carbon.

Este modulo solo lee y convierte (fracciones 0-1 a porcentaje); no escribe en la base.
Cualquier inconsistencia (hoja o columna ausente, valor no numerico, banda desordenada)
detiene la importacion con `ArchivoModeloCarbonInvalidoError` en vez de cargar datos a medias.
"""
from __future__ import annotations

import math
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import duckdb
import openpyxl

from domain.entities.modelo_carbon import (
    ComparacionXM,
    CriterioValidez,
    EpisodioValidacion,
    ErrorPorHorizonte,
    HistoricoCarbonMensual,
    InventarioCarbon,
    PaqueteModeloCarbon,
    ProyeccionCarbonMensual,
    SupuestoCarbon,
    TrimestreCarbon,
)
from domain.value_objects.banda import Banda

RUTA_PROYECCION = Path("resultados") / "proyeccion_carbon.xlsx"
RUTA_EVALUACION = Path("resultados") / "evaluacion" / "evaluacion_modelo.xlsx"
RUTA_BASE_NACIONAL = Path("datos") / "base" / "base_mensual_nacional.parquet"
SUPUESTO_PODER_CALORIFICO = "Poder calorífico [MBTU/t]"

Fila = dict[str, Any]


class ArchivoModeloCarbonInvalidoError(ValueError):
    """Un archivo del modelo de carbon falta, esta incompleto o tiene valores invalidos."""


# ------------------------------------------------------------------ utilidades
def _hoja(libro: openpyxl.Workbook, nombre: str, columnas: list[str], archivo: str) -> list[Fila]:
    if nombre not in libro.sheetnames:
        raise ArchivoModeloCarbonInvalidoError(f"{archivo}: falta la hoja '{nombre}'")
    filas = list(libro[nombre].iter_rows(values_only=True))
    if not filas:
        raise ArchivoModeloCarbonInvalidoError(f"{archivo}: la hoja '{nombre}' esta vacia")
    encabezado = [str(c) if c is not None else "" for c in filas[0]]
    faltan = [c for c in columnas if c not in encabezado]
    if faltan:
        raise ArchivoModeloCarbonInvalidoError(
            f"{archivo}, hoja '{nombre}': faltan columnas {', '.join(faltan)}"
        )
    datos = [dict(zip(encabezado, f)) for f in filas[1:] if any(c is not None for c in f)]
    if not datos:
        raise ArchivoModeloCarbonInvalidoError(f"{archivo}: la hoja '{nombre}' no tiene filas")
    return datos


def _num(valor: Any, donde: str, opcional: bool = False) -> Optional[float]:
    if valor is None or (isinstance(valor, float) and math.isnan(valor)):
        if opcional:
            return None
        raise ArchivoModeloCarbonInvalidoError(f"{donde}: falta un valor numerico")
    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
        raise ArchivoModeloCarbonInvalidoError(f"{donde}: no es numerico ({valor!r})")
    return float(valor)


def _fecha(valor: Any, donde: str) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    raise ArchivoModeloCarbonInvalidoError(f"{donde}: fecha no valida ({valor!r})")


def _pct(valor: Any, donde: str, opcional: bool = False) -> Optional[float]:
    """Fraccion 0-1 del archivo a porcentaje 0-100."""
    numero = _num(valor, donde, opcional)
    return None if numero is None else round(numero * 100, 4)


def _banda(f: Fila, base: str, donde: str, pct: bool = False) -> Banda:
    convertir = _pct if pct else _num
    try:
        return Banda(*(convertir(f[f"{base}_{p}"], f"{donde} {base}_{p}") for p in ("p10", "p50", "p90")))
    except ValueError as error:
        raise ArchivoModeloCarbonInvalidoError(f"{donde}: {error}") from error


def _entero(valor: Any, donde: str) -> int:
    numero = _num(valor, donde)
    assert numero is not None
    return int(numero)


# ------------------------------------------------------------------ conversiones
def _proyeccion(filas: list[Fila]) -> list[ProyeccionCarbonMensual]:
    salida = []
    for i, f in enumerate(filas, start=2):
        donde = f"proyeccion_mensual fila {i}"
        salida.append(
            ProyeccionCarbonMensual(
                fecha=_fecha(f["fecha"], donde), dias=_entero(f["dias"], donde),
                techo_gbtu_dia=_num(f["techo_gbtu_dia"], donde),
                roni=_banda(f, "roni", donde),
                aportes_pct_media=_banda(f, "aportes_pct_media", donde, pct=True),
                uso_techo_pct=_banda(f, "uso_techo", donde, pct=True),
                carbon_gbtu_dia=_banda(f, "carbon_gbtu_dia", donde),
                carbon_kt_mes=_banda(f, "carbon_kt_mes", donde),
                prob_supera_maximo_pct=_pct(f["prob_supera_maximo_historico"], donde),
                observado_parcial_gbtu_dia=_num(f["carbon_observado_parcial"], donde, opcional=True),
            )
        )
    return sorted(salida, key=lambda m: m.fecha)


def _comparacion(filas: list[Fila]) -> list[ComparacionXM]:
    salida = []
    for i, f in enumerate(filas, start=2):
        donde = f"comparacion_xm fila {i}"
        historicas = [
            v for v in (
                _num(f.get(c), f"{donde} {c}", opcional=True)
                for c in ("xm_h15_17_oefatraso", "xm_h23_25_oefatraso", "xm_h91_93_oefatraso")
            ) if v is not None
        ]
        salida.append(
            ComparacionXM(
                fecha=_fecha(f["fecha"], donde),
                modelo=_banda(f, "carbon_gbtu_dia", donde),
                techo_gbtu_dia=_num(f["techo_gbtu_dia"], donde),
                xm_estocastico_gbtu_dia=_num(f["xm_promedio"], donde, opcional=True),
                xm_historicas_min_gbtu_dia=min(historicas) if historicas else None,
                xm_historicas_max_gbtu_dia=max(historicas) if historicas else None,
            )
        )
    return sorted(salida, key=lambda c: c.fecha)


def _historico(ruta: Path) -> list[HistoricoCarbonMensual]:
    if not ruta.is_file():
        raise ArchivoModeloCarbonInvalidoError(f"No existe {ruta}")
    columnas = ["fecha", "carbon_gbtu_dia", "techo_carbon_gbtu_dia", "uso_techo_carbon",
                "aportes_pct_media", "roni", "fuente_carbon", "mes_completo", "dias_con_dato"]
    conexion = duckdb.connect()
    try:
        filas = conexion.execute(
            f"SELECT {', '.join(columnas)} FROM read_parquet(?) ORDER BY fecha", [str(ruta)]
        ).fetchall()
    except duckdb.Error as error:
        raise ArchivoModeloCarbonInvalidoError(f"{ruta.name}: {error}") from error
    finally:
        conexion.close()
    salida = []
    for i, f in enumerate(filas, start=1):
        donde = f"base_mensual_nacional fila {i}"
        if _num(f[1], donde, opcional=True) is None:
            continue  # mes sin consumo calculable
        salida.append(
            HistoricoCarbonMensual(
                fecha=_fecha(f[0], donde), carbon_gbtu_dia=_num(f[1], donde),
                techo_gbtu_dia=_num(f[2], donde, opcional=True), uso_techo_pct=_pct(f[3], donde, True),
                aportes_pct_media=_pct(f[4], donde, True), roni=_num(f[5], donde, opcional=True),
                fuente=str(f[6]), mes_completo=bool(f[7]), dias_con_dato=_entero(f[8], donde),
            )
        )
    return salida


# ------------------------------------------------------------------ punto de entrada
def leer_modelo_carbon(carpeta: Path) -> PaqueteModeloCarbon:
    """Lee la corrida completa del modelo desde la carpeta de su proyecto."""
    ruta_proyeccion, ruta_evaluacion = carpeta / RUTA_PROYECCION, carpeta / RUTA_EVALUACION
    for ruta in (ruta_proyeccion, ruta_evaluacion):
        if not ruta.is_file():
            raise ArchivoModeloCarbonInvalidoError(f"No existe {ruta}")

    libro = openpyxl.load_workbook(ruta_proyeccion, read_only=True, data_only=True)
    evaluacion = openpyxl.load_workbook(ruta_evaluacion, read_only=True, data_only=True)
    try:
        nombre = ruta_proyeccion.name
        proyeccion = _proyeccion(_hoja(libro, "proyeccion_mensual", [
            "fecha", "dias", "techo_gbtu_dia", "roni_p10", "roni_p50", "roni_p90",
            "aportes_pct_media_p10", "aportes_pct_media_p50", "aportes_pct_media_p90",
            "uso_techo_p10", "uso_techo_p50", "uso_techo_p90",
            "carbon_gbtu_dia_p10", "carbon_gbtu_dia_p50", "carbon_gbtu_dia_p90",
            "carbon_kt_mes_p10", "carbon_kt_mes_p50", "carbon_kt_mes_p90",
            "prob_supera_maximo_historico", "carbon_observado_parcial"], nombre))
        trimestres = [
            TrimestreCarbon(str(f["trimestre"]), _entero(f["meses_incluidos"], "trimestres"),
                            _banda(f, "carbon_mt", f"trimestres {f['trimestre']}"))
            for f in _hoja(libro, "trimestres", ["trimestre", "meses_incluidos", "carbon_mt_p10",
                                                 "carbon_mt_p50", "carbon_mt_p90"], nombre)
        ]
        inventario = [
            InventarioCarbon(_fecha(f["fecha"], "inventario"), _entero(f["dias_inventario"], "inventario"),
                             _num(f["inventario_kt_p50"], "inventario"), _num(f["inventario_kt_p90"], "inventario"))
            for f in _hoja(libro, "inventario", ["fecha", "dias_inventario", "inventario_kt_p50",
                                                  "inventario_kt_p90"], nombre)
        ]
        comparacion = _comparacion(_hoja(libro, "comparacion_xm", [
            "fecha", "carbon_gbtu_dia_p10", "carbon_gbtu_dia_p50", "carbon_gbtu_dia_p90",
            "techo_gbtu_dia", "xm_promedio"], nombre))
        episodios = [
            EpisodioValidacion(
                str(f["episodio"]), _fecha(f["inicio"], "validacion"), _fecha(f["fin"], "validacion"),
                _entero(f["meses"], "validacion"), _num(f["carbon_real_promedio"], "validacion"),
                _num(f["carbon_p50_promedio"], "validacion"), _pct(f["error_total_episodio_pct"], "validacion"),
                _num(f["mae_modelo"], "validacion"), _num(f["mae_ingenuo"], "validacion"),
                _pct(f["cobertura_p10_p90"], "validacion"),
            )
            for f in _hoja(libro, "validacion_metricas", [
                "episodio", "inicio", "fin", "meses", "carbon_real_promedio", "carbon_p50_promedio",
                "error_total_episodio_pct", "mae_modelo", "mae_ingenuo", "cobertura_p10_p90"], nombre)
        ]
        supuestos = [
            SupuestoCarbon(str(f["supuesto"]), "" if f["valor"] is None else str(f["valor"]))
            for f in _hoja(libro, "supuestos", ["supuesto", "valor"], nombre)
        ]
        nombre_evaluacion = ruta_evaluacion.name
        veredicto = [
            CriterioValidez(str(f["criterio"]), str(f["resultado"]), str(f["cumple"]).strip().lower() in ("sí", "si", "true"),
                            str(f["interpretacion"]))
            for f in _hoja(evaluacion, "veredicto", ["criterio", "resultado", "cumple", "interpretacion"], nombre_evaluacion)
        ]
        errores = [
            ErrorPorHorizonte(
                _entero(f["horizonte"], "error_por_horizonte"), _entero(f["n"], "error_por_horizonte"),
                _num(f["mae_modelo"], "error_por_horizonte"), _num(f["mae_persistencia"], "error_por_horizonte"),
                _num(f["mae_estacional"], "error_por_horizonte"), _num(f["mae_climatologia"], "error_por_horizonte"),
                _num(f["sesgo_modelo"], "error_por_horizonte"),
            )
            for f in _hoja(evaluacion, "error_por_horizonte", [
                "horizonte", "n", "mae_modelo", "mae_persistencia", "mae_estacional", "mae_climatologia",
                "sesgo_modelo"], nombre_evaluacion)
        ]
    finally:
        libro.close()
        evaluacion.close()

    valores = {s.nombre: s.valor for s in supuestos}
    if SUPUESTO_PODER_CALORIFICO not in valores:
        raise ArchivoModeloCarbonInvalidoError(f"proyeccion_carbon.xlsx: falta el supuesto '{SUPUESTO_PODER_CALORIFICO}'")
    try:
        poder_calorifico = float(valores[SUPUESTO_PODER_CALORIFICO])
    except ValueError as error:
        raise ArchivoModeloCarbonInvalidoError("El poder calorifico del modelo no es numerico") from error

    origen = (
        f"Modelo de carbón térmico (ENOS → aportes → carbón; {valores.get('Simulaciones', '?')} simulaciones "
        f"Monte Carlo) · muestra {valores.get('Muestra de estimación', '?')} · "
        f"pronóstico NOAA del {valores.get('Fecha de emisión pronóstico NOAA', '?')}"
    )
    try:
        return PaqueteModeloCarbon(
            origen=origen, poder_calorifico_mbtu_t=poder_calorifico,
            historico=_historico(carpeta / RUTA_BASE_NACIONAL), proyeccion=proyeccion, trimestres=trimestres,
            inventario=inventario, comparacion_xm=comparacion, errores_horizonte=errores, episodios=episodios,
            veredicto=veredicto, supuestos=supuestos,
        )
    except ValueError as error:
        raise ArchivoModeloCarbonInvalidoError(str(error)) from error
